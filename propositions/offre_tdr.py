"""Offre technique et offre financière en réponse à des termes de référence (TDR).

Produit deux documents Word modifiables à partir d'un fichier de description de l'offre (voir offre_exemple.json).
Les champs entre crochets, comme [NIF], sont surlignés en jaune : ils restent à compléter avant l'envoi.
Usage : python offre_tdr.py <offre.json> <dossier_sortie>
"""
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

O = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
SORTIE = Path(sys.argv[2])
SORTIE.mkdir(parents=True, exist_ok=True)

BLEU = RGBColor(0x18, 0x4F, 0x95)
GRIS = RGBColor(0x55, 0x5F, 0x6B)
FOND_ENTETE, FOND_LEGER, FOND_ACTIF = "184F95", "EAF1FB", "6DA7EC"
MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
        "novembre", "décembre"]

A, S = O["appel"], O["soumissionnaire"]
EQUIPE = {m["id"]: m for m in O["equipe"]}


def jour(iso):
    d = date.fromisoformat(iso) if isinstance(iso, str) else iso
    return f"{'1er' if d.day == 1 else d.day} {MOIS[d.month - 1]} {d.year}"


def ar(v):
    return f"{round(v):,}".replace(",", " ") + " Ar"


# ---------- Montants en lettres ----------
UNITES = ["zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf", "dix", "onze", "douze",
          "treize", "quatorze", "quinze", "seize", "dix-sept", "dix-huit", "dix-neuf"]
DIZAINES = {2: "vingt", 3: "trente", 4: "quarante", 5: "cinquante", 6: "soixante"}


def moins_de_cent(n, final):
    if n < 20:
        return UNITES[n]
    d, u = divmod(n, 10)
    if d == 7:
        return "soixante et onze" if u == 1 else "soixante-" + UNITES[10 + u]
    if d == 8:
        return ("quatre-vingts" if final else "quatre-vingt") if u == 0 else "quatre-vingt-" + UNITES[u]
    if d == 9:
        return "quatre-vingt-" + UNITES[10 + u]
    if u == 0:
        return DIZAINES[d]
    return DIZAINES[d] + (" et un" if u == 1 else "-" + UNITES[u])


def moins_de_mille(n, final=True):
    c, r = divmod(n, 100)
    mots = []
    if c:
        mots.append("cent" if c == 1 else UNITES[c] + " cent" + ("s" if r == 0 and final else ""))
    if r:
        mots.append(moins_de_cent(r, final))
    return " ".join(mots)


def en_lettres(n):
    n = round(n)
    if n == 0:
        return "zéro"
    mots = []
    for valeur, nom in ((10 ** 9, "milliard"), (10 ** 6, "million")):
        q, n = divmod(n, valeur)
        if q:
            mots.append(f"{moins_de_mille(q)} {nom}{'s' if q > 1 else ''}")
    q, n = divmod(n, 1000)
    if q:
        mots.append("mille" if q == 1 else moins_de_mille(q, final=False) + " mille")
    if n:
        mots.append(moins_de_mille(n))
    return " ".join(mots)


# ---------- Calculs ----------
JOURS = {i: sum(e["jours"].get(i, 0) for e in O["methode"]) for i in EQUIPE}
HONORAIRES = sum(JOURS[i] * EQUIPE[i]["tarif"] for i in EQUIPE)
FRAIS = sum(q * p for _, _, q, p in O["frais"])
HT = HONORAIRES + FRAIS
TVA = round(HT * O["tva"]) if O["tva"] else 0
TTC = HT + TVA
DEBUT = date.fromisoformat(O["calendrier"]["debut"])


def semaine(n):
    """Date du vendredi de la semaine n (1 = semaine du démarrage)."""
    return DEBUT + timedelta(days=7 * (n - 1) + 4)


# ---------- Mise en forme Word ----------
def ombrer(cellule, couleur):
    tc = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc.append(shd)


def texte(par, contenu, gras=False, couleur=None, taille=None, italique=False):
    """Ajoute du texte ; les passages entre crochets sont surlignés pour signaler ce qui reste à compléter."""
    for morceau in re.split(r"(\[[^\]]+\])", contenu):
        if not morceau:
            continue
        r = par.add_run(morceau)
        r.bold, r.italic = gras, italique
        if couleur:
            r.font.color.rgb = couleur
        if taille:
            r.font.size = Pt(taille)
        if morceau.startswith("["):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def para(doc, contenu="", **opts):
    alignement = opts.pop("align", None)
    apres = opts.pop("apres", None)
    p = texte(doc.add_paragraph(), contenu, **opts)
    if alignement:
        p.alignment = alignement
    if apres is not None:
        p.paragraph_format.space_after = Pt(apres)
    return p


def puces(doc, elements, style="List Bullet"):
    for e in elements:
        texte(doc.add_paragraph(style=style), e)


def titre(doc, contenu, niveau=1):
    h = doc.add_heading(level=niveau)
    texte(h, contenu, couleur=BLEU)
    return h


def tableau(doc, entetes, lignes, largeurs, total=None):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tr = t.rows[0]._tr.get_or_add_trPr()
    repete = OxmlElement("w:tblHeader")
    repete.set(qn("w:val"), "true")
    tr.append(repete)
    for c, e in zip(t.rows[0].cells, entetes):
        texte(c.paragraphs[0], e, gras=True, couleur=RGBColor(0xFF, 0xFF, 0xFF), taille=9.5)
        ombrer(c, FOND_ENTETE)
    for ligne in lignes:
        cellules = t.add_row().cells
        for c, v in zip(cellules, ligne):
            texte(c.paragraphs[0], str(v), taille=9.5)
            if isinstance(v, str) and v.endswith("Ar") or isinstance(v, (int, float)):
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if total:
        cellules = t.add_row().cells
        for c, v in zip(cellules, total):
            texte(c.paragraphs[0], str(v), gras=True, taille=9.5)
            ombrer(c, FOND_LEGER)
            if str(v).endswith("Ar") or isinstance(v, (int, float)):
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for ligne in t.rows:
        for c, l in zip(ligne.cells, largeurs):
            c.width = Cm(l)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def nouveau(nature):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.paragraph_format.space_after = Pt(5)
    for nom, taille in (("Heading 1", 14), ("Heading 2", 11.5)):
        doc.styles[nom].font.name, doc.styles[nom].font.size = "Calibri", Pt(taille)
    pied = sec.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texte(pied, f"{S['nom']} – {nature} – {O['reference']} – page ", couleur=GRIS, taille=8.5)
    champ = OxmlElement("w:fldSimple")
    champ.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    champ.append(r)
    pied._p.append(champ)
    return doc


def saut(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def garde(doc, nature):
    for _ in range(5):
        doc.add_paragraph()
    para(doc, nature.upper(), gras=True, couleur=BLEU, taille=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, A["intitule"], gras=True, taille=15, align=WD_ALIGN_PARAGRAPH.CENTER, apres=24)
    lignes = [
        ("Appel / TDR", A["reference"]),
        ("Commanditaire", A["commanditaire"]),
        ("Soumissionnaire", S["nom"]),
        ("Référence de l'offre", O["reference"]),
        ("Date", jour(O["date"])),
        ("Date limite de dépôt", jour(A["date_limite"])),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in lignes:
        a, b = t.add_row().cells
        texte(a.paragraphs[0], k, gras=True)
        ombrer(a, FOND_LEGER)
        texte(b.paragraphs[0], v)
        a.width, b.width = Cm(5), Cm(10)
    for _ in range(6):
        doc.add_paragraph()
    para(doc, f"{S['nom']} · {S['email']} · {S['telephone']}", couleur=GRIS, taille=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, S["site"], couleur=GRIS, taille=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    saut(doc)


def entete_lettre(doc, objet):
    para(doc, S["nom"], gras=True, apres=0)
    for l in (S["adresse"], f"NIF : {S['nif']} – STAT : {S['stat']}", f"{S['telephone']} · {S['email']}"):
        para(doc, l, taille=9.5, apres=0)
    doc.add_paragraph()
    for l in (A["destinataire"], A["commanditaire"], A["adresse"]):
        para(doc, l, apres=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_paragraph()
    para(doc, f"[Ville], le {jour(O['date'])}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    p = para(doc, "Objet : ", gras=True)
    texte(p, objet)
    para(doc, f"Référence : {A['reference']}", apres=12)
    para(doc, "Madame, Monsieur,")


def signature(doc):
    doc.add_paragraph()
    para(doc, "Veuillez agréer, Madame, Monsieur, l'expression de nos salutations distinguées.")
    doc.add_paragraph()
    para(doc, S["responsable"], gras=True, apres=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, f"{S['fonction']}, {S['nom']}", apres=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, "[Signature et cachet]", align=WD_ALIGN_PARAGRAPH.RIGHT)


# ---------- Offre technique ----------
def offre_technique():
    doc = nouveau("Offre technique")
    garde(doc, "Offre technique")

    entete_lettre(doc, f"Soumission de l'offre technique – {A['intitule']}")
    para(doc, f"Après avoir pris connaissance des termes de référence {A['reference']}, nous avons l'honneur de vous "
              f"soumettre notre offre technique pour la mission « {A['intitule']} ».")
    para(doc, f"Nous proposons de réaliser la mission en {DUREE} semaines à compter du {jour(DEBUT)}, avec une équipe de "
              f"{len(EQUIPE)} personnes totalisant {sum(JOURS.values())} jours de travail. Notre offre financière est "
              f"présentée dans un document séparé.")
    para(doc, f"La présente offre est valable {O['validite_jours']} jours à compter de la date limite de dépôt. Nous "
              f"certifions que les informations qu'elle contient sont exactes et nous engageons, si elle est retenue, à "
              f"démarrer la mission à la date convenue.")
    signature(doc)
    saut(doc)

    para(doc, "Sommaire", gras=True, couleur=BLEU, taille=14)
    for i, s in enumerate(PARTIES, 1):
        para(doc, f"{i}. {s}", apres=2)
    para(doc, "Annexes : CV de l'équipe, pièces administratives", apres=2, italique=True)
    saut(doc)

    titre(doc, f"1. {PARTIES[0]}")
    for p in S["presentation"]:
        para(doc, p)
    tableau(doc, ["Élément", "Information"], [
        ("Dénomination", S["nom"]), ("Forme juridique", S["statut"]), ("NIF", S["nif"]), ("STAT", S["stat"]),
        ("RCS", S["rcs"]), ("Adresse", S["adresse"]), ("Responsable", S["responsable"]),
        ("Contact", f"{S['telephone']} · {S['email']}"), ("Site", S["site"]),
    ], [5, 11.5])
    titre(doc, "Domaines de compétence", 2)
    puces(doc, S["domaines"])

    titre(doc, f"2. {PARTIES[1]}")
    titre(doc, "Contexte", 2)
    for p in O["contexte"]:
        para(doc, p)
    titre(doc, "Objectifs de la mission", 2)
    puces(doc, O["objectifs"])
    titre(doc, "Notre compréhension du besoin", 2)
    for p in O["comprehension"]:
        para(doc, p)
    titre(doc, "Commentaires et suggestions sur les TDR", 2)
    puces(doc, O["commentaires_tdr"])

    titre(doc, f"3. {PARTIES[2]}")
    para(doc, f"La mission se déroule en {len(O['methode'])} étapes. Chaque étape produit un résultat vérifiable par le "
              f"commanditaire.")
    for i, e in enumerate(O["methode"], 1):
        titre(doc, f"Étape {i}. {e['titre']}", 2)
        para(doc, e["description"])
        p = para(doc, "Résultat : ", gras=True)
        texte(p, e["sortie"])
    titre(doc, "Données utilisées", 2)
    tableau(doc, ["Source", "Usage dans la mission"], O["donnees"], [7.5, 9])
    titre(doc, "Assurance qualité", 2)
    puces(doc, O["qualite"])

    titre(doc, f"4. {PARTIES[3]}")
    para(doc, f"Démarrage le {jour(DEBUT)}, remise de la version finale au plus tard le "
              f"{jour(semaine(DUREE))}. Les cases colorées indiquent les semaines d'activité ; ◆ marque la remise d'un "
              f"livrable.")
    chronogramme(doc)
    titre(doc, "Répartition des jours de travail", 2)
    ids = list(EQUIPE)
    lignes = [[f"{i}. {e['titre']}"] + [e["jours"].get(m, 0) for m in ids] + [sum(e["jours"].values())]
              for i, e in enumerate(O["methode"], 1)]
    tableau(doc, ["Étape"] + [EQUIPE[m]["role"].split(",")[0].split(" (")[0] for m in ids] + ["Total"], lignes,
            [8.5] + [3] * len(ids) + [2], total=["Total"] + [JOURS[m] for m in ids] + [sum(JOURS.values())])

    titre(doc, f"5. {PARTIES[4]}")
    tableau(doc, ["Livrable", "Contenu", "Format", "Remise"],
            [(l["nom"], l["contenu"], l["format"], f"S{l['semaine']} – {jour(semaine(l['semaine']))}")
             for l in O["livrables"]], [3.5, 7, 2.5, 3.5])
    para(doc, "Chaque livrable est remis en version provisoire puis corrigé selon les remarques du commanditaire "
              "(deux tours de corrections inclus).")

    titre(doc, f"6. {PARTIES[5]}")
    tableau(doc, ["Rôle", "Nom", "Profil", "Tâches", "Jours"],
            [(m["role"], m["nom"], m["profil"], m["taches"], JOURS[m["id"]]) for m in O["equipe"]],
            [3.5, 2.8, 4.7, 4.2, 1.3])
    para(doc, "Les CV détaillés et signés figurent en annexe.", italique=True)

    titre(doc, f"7. {PARTIES[6]}")
    for r in O["references"]:
        p = para(doc, "", apres=1)
        texte(p, r["titre"], gras=True)
        texte(p, f" – {r['client']}, {r['annee']}", couleur=GRIS)
        para(doc, r["description"])

    titre(doc, f"8. {PARTIES[7]}")
    tableau(doc, ["Risque", "Mesure prévue"], O["risques"], [7, 9.5])

    titre(doc, f"9. {PARTIES[8]}")
    puces(doc, O["engagements"])

    saut(doc)
    titre(doc, "Annexe 1. CV de l'équipe")
    for m in O["equipe"]:
        titre(doc, f"{m['nom']} – {m['role']}", 2)
        for champ in ("Formation : [Diplômes, établissements, années]",
                      "Expérience : [Postes et missions principales, avec dates et clients]",
                      "Compétences : [Logiciels, méthodes, langues]",
                      "Missions similaires : [Deux ou trois missions comparables]"):
            para(doc, champ)
    titre(doc, "Annexe 2. Pièces administratives à joindre")
    for piece in O["pieces"]:
        texte(doc.add_paragraph(style="List Bullet"), "☐ " + piece)
    nom = SORTIE / f"Offre_technique_{O['reference']}.docx"
    doc.save(nom)
    return nom


def chronogramme(doc):
    n = DUREE
    t = doc.add_table(rows=1, cols=n + 1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    entete = t.rows[0].cells
    texte(entete[0].paragraphs[0], "Étape", gras=True, couleur=RGBColor(0xFF, 0xFF, 0xFF), taille=9)
    ombrer(entete[0], FOND_ENTETE)
    for s in range(1, n + 1):
        c = entete[s]
        texte(c.paragraphs[0], f"S{s}", gras=True, couleur=RGBColor(0xFF, 0xFF, 0xFF), taille=9)
        texte(c.add_paragraph(), f"{semaine(s) - timedelta(days=4):%d/%m}", couleur=RGBColor(0xFF, 0xFF, 0xFF), taille=7)
        ombrer(c, FOND_ENTETE)
    for i, (e, (d, f)) in enumerate(zip(O["methode"], O["calendrier"]["activites"]), 1):
        cellules = t.add_row().cells
        texte(cellules[0].paragraphs[0], f"{i}. {e['titre']}", taille=9)
        for s in range(d, f + 1):
            ombrer(cellules[s], FOND_ACTIF)
    cellules = t.add_row().cells
    texte(cellules[0].paragraphs[0], "Livrables", gras=True, taille=9)
    for l in O["livrables"]:
        c = cellules[l["semaine"]]
        texte(c.paragraphs[0], "◆", gras=True, couleur=BLEU)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ligne in t.rows:
        ligne.cells[0].width = Cm(6.5)
        for c in ligne.cells[1:]:
            c.width = Cm(10 / n)
    doc.add_paragraph()


# ---------- Offre financière ----------
def offre_financiere():
    doc = nouveau("Offre financière")
    garde(doc, "Offre financière")

    entete_lettre(doc, f"Soumission de l'offre financière – {A['intitule']}")
    para(doc, f"Suite à notre offre technique, nous avons l'honneur de vous soumettre notre offre financière pour la "
              f"mission « {A['intitule']} », en réponse aux termes de référence {A['reference']}.")
    p = para(doc, f"Le montant total de notre offre s'élève à ")
    texte(p, f"{en_lettres(TTC)} ariary ({ar(TTC)})", gras=True)
    texte(p, f", {MENTION_TVA}. Ce montant comprend les honoraires et les frais remboursables détaillés ci-après.")
    para(doc, f"La présente offre est valable {O['validite_jours']} jours à compter de la date limite de dépôt.")
    signature(doc)
    saut(doc)

    titre(doc, "1. Récapitulatif")
    lignes = [("Honoraires", ar(HONORAIRES)), ("Frais remboursables", ar(FRAIS)), ("Total hors taxes", ar(HT))]
    lignes.append((f"TVA ({round(O['tva'] * 100)} %)", ar(TVA)) if O["tva"]
                  else ("TVA", "[Non assujetti – à vérifier]"))
    tableau(doc, ["Rubrique", "Montant"], lignes, [11, 5.5], total=["Montant total de l'offre", ar(TTC)])
    p = para(doc, "Arrêté la présente offre à la somme de ")
    texte(p, f"{en_lettres(TTC)} ariary", gras=True)
    texte(p, f" ({ar(TTC)}).")

    titre(doc, "2. Honoraires")
    tableau(doc, ["Membre de l'équipe", "Nom", "Jours", "Tarif journalier", "Montant"],
            [(m["role"], m["nom"], JOURS[m["id"]], ar(m["tarif"]), ar(JOURS[m["id"]] * m["tarif"]))
             for m in O["equipe"]], [5.5, 3.3, 1.5, 3, 3.2],
            total=["Total des honoraires", "", sum(JOURS.values()), "", ar(HONORAIRES)])
    titre(doc, "Honoraires par étape", 2)
    lignes = []
    for i, e in enumerate(O["methode"], 1):
        montant = sum(j * EQUIPE[m]["tarif"] for m, j in e["jours"].items())
        lignes.append((f"{i}. {e['titre']}", sum(e["jours"].values()), ar(montant)))
    tableau(doc, ["Étape", "Jours", "Montant"], lignes, [11, 2, 3.5],
            total=["Total", sum(JOURS.values()), ar(HONORAIRES)])

    titre(doc, "3. Frais remboursables")
    tableau(doc, ["Désignation", "Unité", "Quantité", "Prix unitaire", "Montant"],
            [(l, u, q, ar(p), ar(q * p)) for l, u, q, p in O["frais"]], [6.5, 2, 2, 3, 3],
            total=["Total des frais", "", "", "", ar(FRAIS)])
    para(doc, "Les frais sont facturés sur la base réelle, dans la limite des montants ci-dessus, sur présentation des "
              "justificatifs.", italique=True)

    titre(doc, "4. Échéancier de paiement")
    tableau(doc, ["Tranche", "Condition de paiement", "Part", "Montant"],
            [(i, c, f"{p} %", ar(TTC * p / 100)) for i, (c, p) in enumerate(O["paiement"], 1)], [2, 8.5, 2, 4],
            total=["", "Total", "100 %", ar(TTC)])
    para(doc, "Chaque tranche est payable dans les 30 jours suivant la réception de la facture correspondante.")

    titre(doc, "5. Conditions")
    puces(doc, [
        "Les prix sont fermes et non révisables pendant la durée de la mission.",
        f"Montants exprimés en ariary, {MENTION_TVA}.",
        "Toute prestation supplémentaire demandée par écrit fera l'objet d'un avenant chiffré au même tarif journalier.",
        "Deux tours de corrections sont inclus ; au-delà, les reprises sont facturées au temps passé.",
    ])
    titre(doc, "6. Coordonnées de paiement")
    b = O["banque"]
    tableau(doc, ["Élément", "Information"],
            [("Titulaire", b["titulaire"]), ("Banque", b["banque"]), ("RIB", b["rib"]), ("Mobile money", b["mobile"]),
             ("NIF / STAT", f"{S['nif']} / {S['stat']}")], [5, 11.5])
    nom = SORTIE / f"Offre_financiere_{O['reference']}.docx"
    doc.save(nom)
    return nom


DUREE = O["calendrier"]["semaines"]
PARTIES = ["Présentation du soumissionnaire", "Compréhension des termes de référence", "Méthodologie",
           "Plan de travail et calendrier", "Livrables", "Équipe proposée", "Références", "Risques et mesures",
           "Engagements"]
MENTION_TVA = (f"TVA de {round(O['tva'] * 100)} % comprise" if O["tva"]
               else "hors TVA, le soumissionnaire n'étant pas assujetti [à vérifier]")

if __name__ == "__main__":
    for f in (offre_technique(), offre_financiere()):
        print(f)
    print(f"Honoraires {ar(HONORAIRES)} · frais {ar(FRAIS)} · total {ar(TTC)} ({en_lettres(TTC)} ariary)")
