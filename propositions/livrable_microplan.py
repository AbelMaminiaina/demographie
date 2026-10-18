"""Livrable de démonstration : microplanification d'une campagne de vaccination, district d'Antsirabe II.

Produit un rapport PDF (cartes, tableaux, plan des équipes) et un classeur Excel (communes, fokontany,
équipes, calendrier jour par jour) à partir des données publiées par le site.
Usage : python livrable_microplan.py <dossier_donnees> <dossier_sortie> [--public]
--public : version publiée sur le site (client fictif, aperçus des cartes conservés).
"""
import json
import math
import sys
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Polygon as PolyPatch
from matplotlib.path import Path as MplPath
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

DONNEES, SORTIE = Path(sys.argv[1]), Path(sys.argv[2])
SORTIE.mkdir(parents=True, exist_ok=True)
IMAGES = SORTIE / "_images"
IMAGES.mkdir(exist_ok=True)

DISTRICT = "Antsirabe II"
PUBLIC = "--public" in sys.argv[3:]
CLIENT = "client fictif, exemple de démonstration" if PUBLIC else "ONG Santé Plus Vakinankaratra (exemple)"
REFERENCE = "EXEMPLE" if PUBLIC else "DEV-2026-1018-01"
REF_DEVIS = "exemple public" if PUBLIC else f"devis {REFERENCE}"
# Normes d'exemple, à remplacer par celles du client
PART_CIBLE, PAR_JOUR, JOURS, PERTE = 0.17, 120, 4, 0.10
CAP = PAR_JOUR * JOURS
LIBELLE_CIBLE = "enfants de 0 à 59 mois"
SEUIL_ECART = 30  # % d'écart entre les deux estimations au-delà duquel un fokontany est à vérifier
TYPES_SANTE = ["Centre de santé", "Poste de santé", "Hôpital"]
COUL_SANTE = ["#1baf7a", "#eda100", "#eb6834"]
RAMPE = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#184f95"]


def nb(v):
    return f"{round(v):,}".replace(",", " ")


def dec(v):
    return f"{v:.1f}".replace(".", ",")


# ---------- Données ----------
base = json.loads((DONNEES / "base.json").read_text(encoding="utf-8"))["niveaux"]
F = json.loads((DONNEES / "fokontany.json").read_text(encoding="utf-8"))["unites"]
GF = json.loads((DONNEES / "contours_fokontany.json").read_text(encoding="utf-8"))
GC = {f["properties"]["i"]: f["geometry"] for f in json.loads((DONNEES / "contours_communes.json").read_text(encoding="utf-8"))["features"]}
SANTE = json.loads((DONNEES / "sante.json").read_text(encoding="utf-8"))["points"]
E = GF["echelle"]


def anneaux_fk(i):
    """Contours d'un fokontany : entiers en 1/E de degré, chaque point codé par son écart au précédent."""
    res = []
    for plat in GF["anneaux"][i]:
        x = y = 0
        pts = []
        for k in range(0, len(plat), 2):
            x += plat[k]
            y += plat[k + 1]
            pts.append((x / E, y / E))
        res.append(pts)
    return res


def anneaux_commune(i):
    g = GC[i]
    polys = g["coordinates"] if g["type"] == "MultiPolygon" else [g["coordinates"]]
    return [p[0] for p in polys]


def centre(anneaux):
    a = max(anneaux, key=len)
    return (sum(p[0] for p in a) / len(a), sum(p[1] for p in a) / len(a))


def km(a, b):
    la1, la2 = math.radians(a[1]), math.radians(b[1])
    d = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin(math.radians(b[0] - a[0]) / 2) ** 2
    return 6371 * 2 * math.asin(math.sqrt(d))


di = next(i for i, d in enumerate(base["districts"]) if d["n"] == DISTRICT)
communes = []
for ci, c in enumerate(base["communes"]):
    if c["di"] != di:
        continue
    fks = []
    for fi, f in enumerate(F):
        if f["c"] != ci:
            continue
        an = anneaux_fk(fi)
        wp = f["wp"][11]
        ca = f.get("ca")
        ecart = round(100 * (wp - ca) / ca) if ca else None
        fks.append({"n": f["n"], "pc": f["pc"], "ca": ca, "wp": wp, "ecart": ecart,
                    "verifier": ecart is not None and abs(ecart) >= SEUIL_ECART,
                    "ci": round(wp * PART_CIBLE), "anneaux": an, "centre": centre(an)})
    communes.append({"i": ci, "n": c["n"], "pc": c["pc"], "wp": c["wp"][11], "rg": (c.get("rg") or {}).get("p"),
                     "anneaux": anneaux_commune(ci), "fk": fks})

# Formations sanitaires du district, rattachées à leur fokontany
chemins = [(c, f, [MplPath(a) for a in f["anneaux"]]) for c in communes for f in c["fk"]]
sites = []
for x, y, t, n in SANTE:
    for c, f, ps in chemins:
        if any(p.contains_point((x, y)) for p in ps):
            sites.append({"x": x, "y": y, "t": t, "n": n, "commune": c["n"], "fokontany": f["n"]})
            break
for c in communes:
    for f in c["fk"]:
        proche = min(sites, key=lambda s: km(f["centre"], (s["x"], s["y"])))
        f["site"], f["dist"] = proche, km(f["centre"], (proche["x"], proche["y"]))

# ---------- Plan des équipes ----------
# Un fokontany plus grand que la capacité d'une équipe est partagé en secteurs ; les restes sont regroupés
# avec leurs voisins les plus proches, tant que l'équipe a de la capacité.
def planifier(c):
    equipes, morceaux = [], []
    for f in sorted(c["fk"], key=lambda f: -f["ci"]):
        pleins, reste = divmod(f["ci"], CAP)
        n_sect = pleins + (1 if reste else 0)
        for k in range(pleins):
            equipes.append({"items": [(f"{f['n']} (secteur {k + 1})" if n_sect > 1 else f["n"], CAP, f)]})
        if reste:
            morceaux.append((f"{f['n']} (secteur {n_sect})" if n_sect > 1 else f["n"], reste, f))
    morceaux.sort(key=lambda m: -m[1])
    while morceaux:
        graine = morceaux.pop(0)
        eq, charge = [graine], graine[1]
        while True:
            possibles = [m for m in morceaux if charge + m[1] <= CAP]
            if not possibles:
                break
            m = min(possibles, key=lambda m: km(graine[2]["centre"], m[2]["centre"]))
            morceaux.remove(m)
            eq.append(m)
            charge += m[1]
        equipes.append({"items": eq})
    for k, e in enumerate(equipes, 1):
        e["num"] = k
        e["ci"] = sum(m[1] for m in e["items"])
        e["jours"] = e["ci"] / PAR_JOUR
        e["doses"] = math.ceil(e["ci"] * (1 + PERTE))
        principal = max(e["items"], key=lambda m: m[1])[2]
        e["site"], e["dist"] = principal["site"], principal["dist"]
        # Calendrier : chaque jour, l'équipe vaccine au plus PAR_JOUR enfants, dans l'ordre de ses fokontany
        jours, jour, reste_jour = [], 1, PAR_JOUR
        for nom, n, _ in e["items"]:
            while n > 0:
                fait = min(n, reste_jour)
                jours.append((jour, nom, fait))
                n -= fait
                reste_jour -= fait
                if reste_jour == 0:
                    jour, reste_jour = jour + 1, PAR_JOUR
        e["calendrier"] = jours
    return equipes


for c in communes:
    c["ci"] = sum(f["ci"] for f in c["fk"])
    c["equipes"] = planifier(c)
    c["doses"] = sum(e["doses"] for e in c["equipes"])
    c["a_verifier"] = sum(f["verifier"] for f in c["fk"])
    c["sites"] = [s for s in sites if s["commune"] == c["n"]]
communes.sort(key=lambda c: -c["ci"])
tot = {k: sum(c[k] for c in communes) for k in ("ci", "doses")}
tot["equipes"] = sum(len(c["equipes"]) for c in communes)
tot["fk"] = sum(len(c["fk"]) for c in communes)
tot["wp"] = sum(f["wp"] for c in communes for f in c["fk"])
tot["verifier"] = sum(c["a_verifier"] for c in communes)
tot["sans_site"] = sum(1 for c in communes if not c["sites"])

# ---------- Cartes ----------
toutes = sorted(f["ci"] for c in communes for f in c["fk"])
SEUILS = [toutes[int(q * len(toutes))] for q in (0.2, 0.4, 0.6, 0.8)]


def couleur(v):
    return RAMPE[sum(v >= s for s in SEUILS)]


def carte(liste_communes, fichier, etiquettes, taille):
    fig, ax = plt.subplots(figsize=taille, dpi=200)
    lat = sum(f["centre"][1] for c in liste_communes for f in c["fk"]) / sum(len(c["fk"]) for c in liste_communes)
    ax.set_aspect(1 / math.cos(math.radians(lat)))
    for c in liste_communes:
        for f in c["fk"]:
            for a in f["anneaux"]:
                ax.add_patch(PolyPatch(a, closed=True, facecolor=couleur(f["ci"]), edgecolor="white", linewidth=0.4))
        for a in c["anneaux"]:
            ax.add_patch(PolyPatch(a, closed=True, fill=False, edgecolor="#13201d", linewidth=0.9))
    for s in [s for s in sites if any(s["commune"] == c["n"] for c in liste_communes)]:
        ax.plot(s["x"], s["y"], "o", markersize=5, color=COUL_SANTE[s["t"]], markeredgecolor="white", markeredgewidth=0.7, zorder=5)
    if etiquettes == "communes":
        for c in liste_communes:
            x, y = centre(c["anneaux"])
            ax.text(x, y, c["n"].split(" ")[0] if len(c["n"]) > 16 else c["n"], fontsize=6.5, ha="center", va="center",
                    fontweight="bold", color="#13201d", bbox=dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.75))
    else:
        for c in liste_communes:
            for k, f in enumerate(sorted(c["fk"], key=lambda f: -f["ci"]), 1):
                f["repere"] = k
                ax.text(*f["centre"], str(k), fontsize=7, ha="center", va="center", fontweight="bold", color="#13201d",
                        bbox=dict(boxstyle="circle,pad=0.2", fc="white", ec="#13201d", lw=0.5))
    ax.autoscale_view()
    ax.axis("off")
    fig.savefig(fichier, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


carte(communes, IMAGES / "district.png", "communes", (7, 9))
for c in communes:
    carte([c], IMAGES / f"commune_{c['pc']}.png", "fokontany", (6, 5))

# ---------- Rapport PDF ----------
police = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
pdfmetrics.registerFont(TTFont("Texte", str(police / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("Texte-Gras", str(police / "DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("Texte", normal="Texte", bold="Texte-Gras", italic="Texte", boldItalic="Texte-Gras")
BLEU, ENCRE, GRIS, CLAIR = colors.HexColor("#2a78d6"), colors.HexColor("#13201d"), colors.HexColor("#4c5955"), colors.HexColor("#f0f2ef")
S = {
    "titre": ParagraphStyle("titre", fontName="Texte-Gras", fontSize=20, leading=25, textColor=ENCRE),
    "sur": ParagraphStyle("sur", fontName="Texte-Gras", fontSize=8, leading=11, textColor=BLEU),
    "h1": ParagraphStyle("h1", fontName="Texte-Gras", fontSize=15, leading=19, textColor=ENCRE, spaceBefore=4, spaceAfter=6),
    "h2": ParagraphStyle("h2", fontName="Texte-Gras", fontSize=10.5, leading=14, textColor=BLEU, spaceBefore=8, spaceAfter=4),
    "p": ParagraphStyle("p", fontName="Texte", fontSize=9, leading=13, textColor=ENCRE, spaceAfter=5),
    "note": ParagraphStyle("note", fontName="Texte", fontSize=7.8, leading=10.5, textColor=GRIS, spaceAfter=4),
    "puce": ParagraphStyle("puce", fontName="Texte", fontSize=9, leading=13, textColor=ENCRE, leftIndent=10, bulletIndent=0, spaceAfter=2),
    "cel": ParagraphStyle("cel", fontName="Texte", fontSize=7.6, leading=9.5, textColor=ENCRE),
    "chiffre": ParagraphStyle("chiffre", fontName="Texte-Gras", fontSize=17, leading=20, textColor=ENCRE),
    "legende": ParagraphStyle("legende", fontName="Texte", fontSize=7.6, leading=10, textColor=GRIS),
}


def tableau(entete, lignes, largeurs, total=None, droite=(1, -1)):
    data = [entete] + lignes + ([total] if total else [])
    t = Table(data, colWidths=[w * mm for w in largeurs], repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, 0), "Texte-Gras", 6.8), ("TEXTCOLOR", (0, 0), (-1, 0), GRIS), ("BACKGROUND", (0, 0), (-1, 0), CLAIR),
        ("FONT", (0, 1), (-1, -1), "Texte", 7.8), ("ALIGN", (droite[0], 0), (droite[1], -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#dfe3e0")), ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
    ]
    if total:
        style += [("FONT", (0, -1), (-1, -1), "Texte-Gras", 8), ("LINEABOVE", (0, -1), (-1, -1), 0.9, ENCRE)]
    t.setStyle(TableStyle(style))
    return t


def pied(canvas, doc):
    canvas.saveState()
    canvas.setFont("Texte", 7)
    canvas.setFillColor(GRIS)
    canvas.drawString(18 * mm, 10 * mm, f"Démographie de Madagascar · Microplanification, district d'{DISTRICT} · {REFERENCE} · version de démonstration")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))
    canvas.restoreState()


histoire = []
P = lambda t, s="p": histoire.append(Paragraph(t, S[s]))

# Couverture
histoire += [Spacer(1, 30 * mm)]
P("LIVRABLE · MICROPLANIFICATION DE CAMPAGNE", "sur")
histoire.append(Spacer(1, 3 * mm))
P(f"Campagne de vaccination de mars 2027<br/>District d'{DISTRICT}", "titre")
histoire.append(Spacer(1, 6 * mm))
P(f"Population cible, besoins en équipes et en doses, et plan de répartition des équipes pour les {len(communes)} communes et {tot['fk']} fokontany du district.")
P(f"Client : {CLIENT}<br/>Référence : {REF_DEVIS}, formule Standard<br/>Établi le {date.today().strftime('%d/%m/%Y')} par Démographie de Madagascar")
histoire.append(Spacer(1, 8 * mm))
chiffres = [[Paragraph(nb(tot["ci"]), S["chiffre"]), Paragraph(str(tot["equipes"]), S["chiffre"]), Paragraph(nb(tot["doses"]), S["chiffre"]), Paragraph(str(tot["verifier"]), S["chiffre"])],
            [Paragraph(f"{LIBELLE_CIBLE} ciblés", S["legende"]), Paragraph(f"équipes sur {JOURS} jours", S["legende"]),
             Paragraph("doses à prévoir (perte comprise)", S["legende"]), Paragraph("fokontany à vérifier sur place", S["legende"])]]
t = Table(chiffres, colWidths=[43 * mm] * 4)
t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEABOVE", (0, 0), (-1, 0), 1.2, BLEU), ("TOPPADDING", (0, 0), (-1, 0), 6)]))
histoire.append(t)
histoire.append(Spacer(1, 12 * mm))
encadre = Table([[Paragraph("<b>Version de démonstration.</b> Ce livrable applique des normes d'exemple (ci-dessous). "
                            "La version définitive reprend vos normes, votre liste de sites de vaccination et les chiffres de la dernière campagne, "
                            "et les fokontany signalés sont vérifiés avec les chefs fokontany.", S["p"])]], colWidths=[174 * mm])
encadre.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fbf0dc")), ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#a66300")),
                             ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
histoire.append(encadre)
histoire.append(PageBreak())

# Synthèse et méthode
P("1. Synthèse", "h1")
gros = communes[:3]
P(f"Le district compte environ {nb(tot['wp'])} habitants en 2026 (estimation WorldPop), dont {nb(tot['ci'])} {LIBELLE_CIBLE}. "
  f"Avec {PAR_JOUR} enfants vaccinés par équipe et par jour pendant {JOURS} jours, la campagne mobilise <b>{tot['equipes']} équipes</b> "
  f"et <b>{nb(tot['doses'])} doses</b>. Les trois communes les plus chargées ({', '.join(c['n'] for c in gros)}) "
  f"concentrent {round(100 * sum(c['ci'] for c in gros) / tot['ci'])} % de la cible.")
annonce = sum(math.ceil(round(c["wp"] * PART_CIBLE) / CAP) for c in communes)
P(f"La proposition estimait {annonce} équipes à partir des totaux communaux. Le découpage par fokontany en demande "
  f"{tot['equipes']} : une équipe ne peut pas toujours remplir ses {JOURS} jours "
  f"quand les fokontany qu'on peut regrouper sont trop petits ou trop éloignés. C'est l'apport principal de la microplanification : "
  f"ce besoin réel apparaît avant la campagne, pas pendant.")
P(f"<b>Points d'attention.</b> {tot['verifier']} fokontany sur {tot['fk']} présentent un écart de {SEUIL_ECART} % ou plus entre les deux estimations "
  f"de population : leurs effectifs sont à confirmer avec les chefs fokontany avant de figer le plan. "
  + (f"{tot['sans_site']} commune(s) n'ont aucune formation sanitaire publique recensée : la vaccination y passe par des postes avancés ou le porte-à-porte. " if tot["sans_site"] else "")
  + "Les distances aux formations sanitaires sont à vol d'oiseau ; les pistes et la saison des pluies les allongent.")
P("2. Méthode et hypothèses", "h1")
for t_ in [
    "<b>Population 2026</b> : estimation WorldPop sur grille de 100 m, sommée par fokontany. La population 2018 « calée » répartit le recensement RGPH-3 de chaque commune entre ses fokontany.",
    f"<b>Cible</b> : {LIBELLE_CIBLE} = {round(PART_CIBLE * 100)} % de la population (norme d'exemple).",
    f"<b>Rendement</b> : une équipe vaccine {PAR_JOUR} enfants par jour pendant {JOURS} jours, soit {CAP} enfants au plus (norme d'exemple).",
    f"<b>Doses</b> : une par enfant, plus {round(PERTE * 100)} % de perte (norme d'exemple).",
    "<b>Équipes</b> : un fokontany plus grand que la capacité d'une équipe est partagé en secteurs ; les fokontany restants sont regroupés avec leurs voisins les plus proches.",
    "<b>Formations sanitaires</b> : établissements publics recensés ; chaque fokontany est rattaché à la plus proche.",
]:
    histoire.append(Paragraph(t_, S["puce"], bulletText="•"))
P("Sources : recensement RGPH-3 2018 (INSTAT), WorldPop (R2025A), formations sanitaires publiques (HDX), contours geoBoundaries et fokontany (OCHA). "
  "Sous le district, tous les chiffres sont des estimations.", "note")
histoire.append(PageBreak())

# Carte et tableau du district
P("3. Vue d'ensemble du district", "h1")
histoire.append(Image(str(IMAGES / "district.png"), width=120 * mm, height=120 * mm * 9 / 7, kind="proportional"))
leg = "Couleur : " + " · ".join(f'<font color="{RAMPE[k]}">■</font> {nb(a)}–{nb(b)}' for k, (a, b) in enumerate(zip([toutes[0]] + SEUILS, SEUILS + [toutes[-1]])))
P(f"{leg} {LIBELLE_CIBLE} par fokontany. Points : " + " · ".join(f'<font color="{COUL_SANTE[k]}">●</font> {n}' for k, n in enumerate(TYPES_SANTE) if any(s["t"] == k for s in sites))
  + ". Traits noirs : limites des communes. Zone blanche au centre : ville d'Antsirabe (district d'Antsirabe I).", "legende")
histoire.append(PageBreak())
P("4. Besoins par commune", "h1")
lignes = [[c["n"], str(len(c["fk"])), nb(sum(f["wp"] for f in c["fk"])), nb(c["ci"]), str(len(c["equipes"])), nb(c["doses"]), str(len(c["sites"])), str(c["a_verifier"])] for c in communes]
histoire.append(tableau(["COMMUNE", "FOKONTANY", "POP. 2026", "CIBLE", "ÉQUIPES", "DOSES", "FORM. SAN.", "À VÉRIFIER"], lignes,
                        [50, 18, 21, 19, 17, 19, 18, 18], ["Total du district", str(tot["fk"]), nb(tot["wp"]), nb(tot["ci"]), str(tot["equipes"]), nb(tot["doses"]), str(len(sites)), str(tot["verifier"])]))
P("« À vérifier » : nombre de fokontany dont les deux estimations divergent de 30 % ou plus.", "note")

# Une section par commune
for n_c, c in enumerate(communes, 1):
    histoire.append(PageBreak())
    P(f"COMMUNE {n_c} SUR {len(communes)}", "sur")
    P(c["n"], "h1")
    P(f"{len(c['fk'])} fokontany · {nb(sum(f['wp'] for f in c['fk']))} habitants en 2026 · <b>{nb(c['ci'])} {LIBELLE_CIBLE}</b> · "
      f"<b>{len(c['equipes'])} équipes</b> · {nb(c['doses'])} doses · {len(c['sites'])} formation(s) sanitaire(s) publique(s)")
    histoire.append(Image(str(IMAGES / f"commune_{c['pc']}.png"), width=150 * mm, height=95 * mm, kind="proportional"))
    P("Les numéros renvoient au tableau ci-dessous. Couleurs et points : même légende que la carte du district.", "legende")
    P("Population cible par fokontany", "h2")
    lf = []
    for f in sorted(c["fk"], key=lambda f: f["repere"]):
        lf.append([str(f["repere"]), Paragraph(f["n"], S["cel"]), nb(f["ca"]) if f["ca"] else "n.d.", nb(f["wp"]),
                   "" if f["ecart"] is None else f"{f['ecart']:+d} %" + (" ⚠" if f["verifier"] else ""), nb(f["ci"]),
                   Paragraph(f"{f['site']['n']} ({dec(f['dist'])} km)", S["cel"])])
    histoire.append(tableau(["N°", "FOKONTANY", "2018 CALÉE", "2026", "ÉCART", "CIBLE", "FORMATION SANITAIRE LA PLUS PROCHE"], lf,
                            [8, 42, 18, 16, 16, 14, 60], ["", "Total", nb(sum(f["ca"] or 0 for f in c["fk"])), nb(sum(f["wp"] for f in c["fk"])), "", nb(c["ci"]), ""], droite=(2, 5)))
    P(f"⚠ : écart de {SEUIL_ECART} % ou plus entre les deux estimations, effectif à confirmer avec le chef fokontany.", "note")
    le = [[f"Éq. {e['num']}", Paragraph(" + ".join(m[0] for m in e["items"]), S["cel"]), nb(e["ci"]), dec(e["jours"]), nb(e["doses"]),
           Paragraph(e["site"]["n"], S["cel"])] for e in c["equipes"]]
    titre_plan = Paragraph("Plan de répartition des équipes", S["h2"])
    titre_plan.keepWithNext = True
    histoire.append(titre_plan)
    histoire.append(tableau(["ÉQUIPE", "FOKONTANY OU SECTEURS", "CIBLE", "JOURS", "DOSES", "BASE (FORMATION SANITAIRE)"], le,
                            [14, 70, 14, 13, 14, 49], droite=(2, 4)))
    P("Le calendrier jour par jour de chaque équipe figure dans le classeur Excel (onglet « Calendrier »).", "note")

# Limites
histoire.append(PageBreak())
P("5. Limites et prochaines étapes", "h1")
for t_ in [
    "Les populations par fokontany sont des estimations modélisées ; l'écart entre les deux sources dépasse parfois 100 %. Les dénombrements des chefs fokontany font foi.",
    "Les contours des fokontany et des communes ne coïncident pas exactement : les totaux par commune peuvent différer légèrement des chiffres communaux publiés.",
    "Les distances sont à vol d'oiseau. Le regroupement des fokontany suit la proximité, pas encore le réseau de pistes.",
    "Prochaines étapes : intégrer vos normes et votre liste de sites de vaccination, confirmer les fokontany signalés, puis figer le plan et produire les fiches d'équipe imprimables.",
]:
    histoire.append(Paragraph(t_, S["puce"], bulletText="•"))

doc = SimpleDocTemplate(str(SORTIE / f"Rapport_microplanification_{DISTRICT.replace(' ', '_')}.pdf"), pagesize=A4,
                        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm,
                        title=f"Microplanification, district d'{DISTRICT}", author="Démographie de Madagascar")
doc.build(histoire, onFirstPage=pied, onLaterPages=pied)

# ---------- Classeur Excel ----------
wb = Workbook()
GRAS, ENTETE = Font(bold=True), PatternFill("solid", fgColor="F0F2EF")


def feuille(ws, entete, lignes, largeurs):
    ws.append(entete)
    for cell in ws[1]:
        cell.font, cell.fill = GRAS, ENTETE
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for l in lignes:
        ws.append(l)
    for k, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


ws = wb.active
ws.title = "Lisez-moi"
for l in [[f"Microplanification de la campagne de vaccination – district d'{DISTRICT}"], [f"Client : {CLIENT} · {REF_DEVIS}"],
          ["VERSION DE DÉMONSTRATION : normes d'exemple, à remplacer par celles du client."], [],
          ["Normes utilisées"], ["Population cible", f"{LIBELLE_CIBLE} = {round(PART_CIBLE * 100)} % de la population"],
          ["Rendement", f"{PAR_JOUR} enfants par équipe et par jour"], ["Durée de la campagne", f"{JOURS} jours"],
          ["Perte de doses", f"{round(PERTE * 100)} %"], [],
          ["Onglets"], ["Communes", "Besoins par commune : cible, équipes, doses"],
          ["Fokontany", "Population 2018 calée et 2026 WorldPop, écart, cible, formation sanitaire la plus proche"],
          ["Équipes", "Fokontany ou secteurs attribués à chaque équipe, cible, jours, doses"],
          ["Calendrier", "Programme jour par jour de chaque équipe"], [],
          ["Sources", "RGPH-3 2018 (INSTAT), WorldPop R2025A, formations sanitaires publiques (HDX), contours OCHA et geoBoundaries"]]:
    ws.append(l)
ws["A1"].font = Font(bold=True, size=13)
ws["A3"].font = Font(bold=True, color="A66300")
for cell in ("A5", "A11"):
    ws[cell].font = GRAS
ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 90

feuille(wb.create_sheet("Communes"), ["Commune", "Code", "Fokontany", "Population 2026", "Cible", "Équipes", "Doses", "Formations sanitaires", "Fokontany à vérifier"],
        [[c["n"], c["pc"], len(c["fk"]), sum(f["wp"] for f in c["fk"]), c["ci"], len(c["equipes"]), c["doses"], len(c["sites"]), c["a_verifier"]] for c in communes],
        [28, 14, 11, 16, 10, 10, 10, 14, 14])
feuille(wb.create_sheet("Fokontany"), ["Commune", "Fokontany", "Code", "Population 2018 calée", "Population 2026 WorldPop", "Écart (%)", "À vérifier",
                                       "Cible", "Jours-équipe", "Formation sanitaire la plus proche", "Distance (km, vol d'oiseau)"],
        [[c["n"], f["n"], f["pc"], f["ca"], f["wp"], f["ecart"], "oui" if f["verifier"] else "", f["ci"], round(f["ci"] / PAR_JOUR, 1), f["site"]["n"], round(f["dist"], 1)]
         for c in communes for f in sorted(c["fk"], key=lambda f: -f["ci"])],
        [24, 32, 16, 14, 14, 10, 10, 10, 11, 34, 14])
feuille(wb.create_sheet("Équipes"), ["Commune", "Équipe", "Fokontany ou secteurs", "Cible", "Jours", "Doses", "Base (formation sanitaire)", "Distance (km)"],
        [[c["n"], e["num"], " + ".join(m[0] for m in e["items"]), e["ci"], round(e["jours"], 1), e["doses"], e["site"]["n"], round(e["dist"], 1)]
         for c in communes for e in c["equipes"]],
        [24, 8, 60, 10, 8, 10, 34, 12])
feuille(wb.create_sheet("Calendrier"), ["Commune", "Équipe", "Jour", "Fokontany ou secteur", "Enfants à vacciner"],
        [[c["n"], e["num"], j, nom, n] for c in communes for e in c["equipes"] for j, nom, n in e["calendrier"]],
        [24, 8, 6, 44, 16])
wb.save(SORTIE / f"Microplan_{DISTRICT.replace(' ', '_')}.xlsx")

if PUBLIC:
    # Aperçus pour la page du site : la carte du district et celle de la commune la plus peuplée
    grande = max(communes, key=lambda c: c["ci"])
    (IMAGES / "district.png").replace(SORTIE / "apercu_district.png")
    (IMAGES / f"commune_{grande['pc']}.png").replace(SORTIE / "apercu_commune.png")
    print(f"Commune en aperçu : {grande['n']}")
for p in IMAGES.iterdir():
    p.unlink()
IMAGES.rmdir()
print(f"{len(communes)} communes, {tot['fk']} fokontany, {len(sites)} formations sanitaires, {tot['equipes']} équipes, {nb(tot['ci'])} enfants, {tot['verifier']} à vérifier")
